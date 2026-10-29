"""
把中文稿件导出为 Word 文档。

格式要求：
- 主持人提问段落以「Q：」开头，整段加粗
- 选手回答段落以「A：」开头，不加粗
- 全文小四（12pt），字体保持 Word 默认，其余任何地方都不加粗
"""
import re
from pathlib import Path
from typing import List, Tuple

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

QUESTION = re.compile(r"^\s*[Qq]\s*[:：]\s*")
ANSWER = re.compile(r"^\s*[Aa]\s*[:：]\s*")
FONT_SIZE = Pt(12)  # 小四


def parse_paragraphs(text: str) -> List[Tuple[str, str]]:
    """把稿件文本拆成 (kind, text) 段落列表，kind 为 q / a / p。"""
    paragraphs = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if QUESTION.match(line):
            paragraphs.append(("q", "Q：" + QUESTION.sub("", line, count=1)))
        elif ANSWER.match(line):
            paragraphs.append(("a", "A：" + ANSWER.sub("", line, count=1)))
        else:
            paragraphs.append(("p", line))
    return paragraphs


def export_docx(text: str, path: Path) -> Path:
    document = Document()
    section = document.sections[0]  # 模板默认是 Letter，改成 A4 与国内 Word 默认页边距
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(3.18)
    normal = document.styles["Normal"]
    normal.font.size = FONT_SIZE
    normal.font.bold = False
    # 东亚语言标记为简体中文，Word 才会用默认的中文字体渲染
    r_pr = normal.element.get_or_add_rPr()
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = r_pr.makeelement(qn("w:lang"), {})
        r_pr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")

    for kind, content in parse_paragraphs(text):
        run = document.add_paragraph().add_run(content)
        run.font.size = FONT_SIZE
        run.bold = kind == "q"

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path
